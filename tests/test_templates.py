"""Testes dos parâmetros que os templates de contrato declaram."""

import pytest

from contract_generator.models.formatters import (formatar_moeda,
                                                  formatar_percentual)
from contract_generator.models.template import (CampoExtra,
                                                carregar_campos_extras,
                                                ler_template,
                                                listar_tipos_disponiveis)
from contract_generator.models.clause import carregar_clausulas_padrao


def test_tipos_disponiveis_incluem_fotografia():
    assert "fotografia" in listar_tipos_disponiveis()


def test_template_em_formato_lista_continua_funcionando():
    """servico.json é uma lista pura de cláusulas (formato original)."""
    template = ler_template("servico")
    assert template["campos"] == []
    assert len(template["clausulas"]) > 0


def test_template_desconhecido_lista_os_tipos_validos():
    with pytest.raises(FileNotFoundError, match="fotografia"):
        ler_template("inexistente")


def test_fotografia_declara_so_a_duracao():
    """As demais condições não variam por contrato: viraram texto na cláusula."""
    nomes = {campo.nome for campo in carregar_campos_extras("fotografia")}
    assert nomes == {"duracao_horas"}


def test_fotografia_traz_a_descricao_padrao_do_ensaio():
    assert ler_template("fotografia")["descricao_padrao"] == "Ensaio Fotográfico"
    assert ler_template("servico")["descricao_padrao"] == ""


def test_todo_placeholder_das_clausulas_tem_campo_ou_e_do_nucleo():
    """Um placeholder sem origem sairia cru no meio do contrato."""
    import re

    nucleo = {"valor", "valor_metade_1", "valor_metade_2", "forma_pagamento",
              "data_inicio", "data_fim", "descricao_servico", "numero",
              "data_criacao"}
    for tipo in listar_tipos_disponiveis():
        campos = {campo.nome for campo in carregar_campos_extras(tipo)}
        for clausula in carregar_clausulas_padrao(tipo):
            usados = set(re.findall(r"{(\w+)}", clausula.conteudo))
            orfaos = usados - campos - nucleo
            # Os demais placeholders vêm de Contrato.to_dict (partes/endereço).
            assert all(o.startswith(("contratante_", "contratado_"))
                       for o in orfaos), f"{tipo}/{clausula.numero}: {orfaos}"


@pytest.mark.parametrize("tipo, valor, esperado", [
    ("moeda", "450", "R$ 450,00"),
    ("moeda", 1234.5, "R$ 1.234,50"),
    ("percentual", "15", "15%"),
    ("percentual", "12.5", "12,5%"),
    ("inteiro", "20.0", "20"),
    ("texto", "  PIXIESET  ", "PIXIESET"),
    ("numero", "", ""),
])
def test_campo_formata_o_valor_conforme_o_tipo(tipo, valor, esperado):
    campo = CampoExtra(nome="x", label="X", tipo=tipo)
    assert campo.formatar(valor) == esperado


def test_campo_numerico_preserva_texto_invalido():
    """Melhor sair o que o usuário digitou do que engolir o dado."""
    campo = CampoExtra(nome="x", label="X", tipo="moeda")
    assert campo.formatar("a combinar") == "a combinar"


def test_campo_com_tipo_invalido_e_recusado():
    with pytest.raises(ValueError, match="inválido"):
        CampoExtra(nome="x", label="X", tipo="booleano")


def test_formatadores():
    assert formatar_moeda(900) == "R$ 900,00"
    assert formatar_percentual(25) == "25%"


def test_extras_nao_sobrescrevem_campos_do_nucleo():
    from datetime import date
    from contract_generator.models.contract import Contrato
    from contract_generator.models.party import Parte

    parte = Parte("Ana", "12345678909", "CPF", "", "")
    contrato = Contrato(
        tipo="fotografia", numero="1/2026", data_criacao=date(2026, 1, 1),
        data_inicio=date(2026, 1, 1), data_fim=None,
        contratante=parte, contratado=parte,
        clausulas=carregar_clausulas_padrao("fotografia"),
        valor=900.0, forma_pagamento="PIX",
        extras={"valor": "R$ 1,00", "duracao_horas": "3"},
    )
    dados = contrato.to_dict()
    assert dados["valor"] == "R$ 900,00"
    assert dados["duracao_horas"] == "3"


# --- Integração com o formulário -------------------------------------------

import io  # noqa: E402

import docx  # noqa: E402

from .conftest import registrar  # noqa: E402

DADOS_BASE = {
    "contract_type": "fotografia", "value": "900",
    "payment_method": "PIX", "start_date": "2026-08-21", "end_date": "",
    "description": "Ensaio externo",
    "contratante_nome": "Maria", "contratante_tipo_documento": "CPF",
    "contratante_documento": "12345678909", "contratante_cidade": "Florianópolis",
    "contratado_nome": "Ana", "contratado_tipo_documento": "CPF",
    "contratado_documento": "98765432100",
}


def _texto_do_docx(app):
    from app.models import ContractRecord

    with app.app_context():
        record = ContractRecord.query.first()
        documento = docx.Document(io.BytesIO(record.docx_data))
        return "\n".join(p.text for p in documento.paragraphs)


def test_formulario_mostra_os_campos_do_template(client):
    registrar(client)
    pagina = client.get("/contracts/new").get_data(as_text=True)
    assert 'data-tipo="fotografia"' in pagina
    assert 'name="extra_duracao_horas"' in pagina
    assert 'value="2"' in pagina  # padrão declarado no template


def test_formulario_entrega_as_descricoes_padrao_ao_navegador(client):
    registrar(client)
    pagina = client.get("/contracts/new").get_data(as_text=True)
    assert 'id="descricoes-padrao"' in pagina
    assert "Ensaio Fotogr" in pagina


def test_formulario_oferece_os_meios_de_pagamento(client):
    registrar(client)
    pagina = client.get("/contracts/new").get_data(as_text=True)
    assert 'id="pagamento-meios"' in pagina
    assert "Transferência bancária" in pagina


def test_parametros_informados_chegam_ao_contrato(client, app):
    registrar(client)
    client.post("/contracts/new", data=dict(DADOS_BASE, extra_duracao_horas="3"),
                follow_redirects=True)

    texto = _texto_do_docx(app)
    assert "duração de 3 horas" in texto


def test_campos_omitidos_caem_no_padrao_do_template(client, app):
    registrar(client)
    client.post("/contracts/new", data=DADOS_BASE, follow_redirects=True)

    texto = _texto_do_docx(app)
    assert "duração de 2 horas" in texto
    assert "{" not in texto  # nenhum placeholder cru sobrou


def test_condicoes_fixas_saem_cravadas_no_texto(client, app):
    """Prazos, multas, plataforma e foro deixaram de ser parametrizáveis."""
    registrar(client)
    client.post("/contracts/new", data=DADOS_BASE, follow_redirects=True)

    texto = _texto_do_docx(app)
    assert "20 dias úteis" in texto
    assert "via PIXIESET" in texto
    assert "multa no valor de 25%" in texto
    assert "comarca de Florianópolis" in texto


def test_preco_sai_dividido_em_duas_metades(client, app):
    registrar(client)
    client.post("/contracts/new", data=DADOS_BASE, follow_redirects=True)

    texto = _texto_do_docx(app)
    assert "importância de R$ 900,00" in texto
    assert "R$ 450,00 no dia da assinatura" in texto
    assert "R$ 450,00 no dia do ensaio" in texto


# --- Metades do pagamento ---------------------------------------------------

@pytest.mark.parametrize("valor, esperado", [
    (900.0, ("R$ 450,00", "R$ 450,00")),
    (901.01, ("R$ 450,51", "R$ 450,50")),   # o centavo ímpar vai na primeira
    (0.01, ("R$ 0,01", "R$ 0,00")),
    (1234.57, ("R$ 617,29", "R$ 617,28")),
])
def test_metades_do_pagamento_fecham_com_o_total(valor, esperado):
    from datetime import date
    from contract_generator.models.contract import Contrato
    from contract_generator.models.party import Parte

    parte = Parte("Ana", "12345678909", "CPF", "", "")
    contrato = Contrato(
        tipo="fotografia", data_criacao=date(2026, 1, 1),
        data_inicio=date(2026, 1, 1), data_fim=None,
        contratante=parte, contratado=parte,
        clausulas=carregar_clausulas_padrao("fotografia"),
        valor=valor, forma_pagamento="PIX",
    )
    dados = contrato.to_dict()
    assert (dados["valor_metade_1"], dados["valor_metade_2"]) == esperado

    metade_1, metade_2 = contrato.metades_do_pagamento()
    assert round(metade_1 + metade_2, 2) == valor


def test_metades_nao_podem_ser_sobrescritas_por_um_template():
    """São chaves do núcleo: os extras não têm precedência sobre elas."""
    from datetime import date
    from contract_generator.models.contract import Contrato
    from contract_generator.models.party import Parte

    parte = Parte("Ana", "12345678909", "CPF", "", "")
    contrato = Contrato(
        tipo="fotografia", data_criacao=date(2026, 1, 1),
        data_inicio=date(2026, 1, 1), data_fim=None,
        contratante=parte, contratado=parte,
        clausulas=carregar_clausulas_padrao("fotografia"),
        valor=900.0, forma_pagamento="PIX",
        extras={"valor_metade_1": "R$ 1,00"},
    )
    assert contrato.to_dict()["valor_metade_1"] == "R$ 450,00"


def test_cpf_invalido_no_formulario_nao_gera_contrato(client, app):
    """POST com CPF inválido volta com erro no campo e não salva nada."""
    from app.models import ContractRecord

    registrar(client)
    dados = dict(DADOS_BASE, contratante_documento="12345678901")
    resp = client.post("/contracts/new", data=dados, follow_redirects=True)
    html = resp.get_data(as_text=True)

    assert "CPF inválido" in html
    with app.app_context():
        assert ContractRecord.query.count() == 0


def test_vigencia_invertida_no_formulario_e_reportada(client, app):
    registrar(client)
    dados = dict(DADOS_BASE, start_date="2026-12-01", end_date="2026-01-01")
    resp = client.post("/contracts/new", data=dados, follow_redirects=True)
    assert "anterior à data de início" in resp.get_data(as_text=True)


